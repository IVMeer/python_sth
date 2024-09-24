import openpyxl
from docx import Document


# 读取 Excel 文件
def read_excel(file_path):
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)
    return data


# 将数据写入 Word 文件
def write_to_word(data, output_path):
    doc = Document()

    # 遍历 Excel 数据并写入到 Word 文档中
    for row in data:
        row_text = ' | '.join([str(cell) for cell in row if cell is not None])
        doc.add_paragraph(row_text)  # 每行 Excel 数据作为一个段落

    doc.save(output_path)


# 主函数
def excel_to_word(excel_path, word_path):
    data = read_excel(excel_path)
    write_to_word(data, word_path)
    print(f"数据已从 {excel_path} 转换为 {word_path}")


# 示例：将 Excel 文件转换为 Word 文件
excel_to_word('C:/Users/win11/Desktop/123.xlsx', 'C:/Users/win11/Desktop/111.doc')
